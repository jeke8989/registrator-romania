from datetime import datetime
from pprint import pprint
import re
import dateutil
import dateutil.parser
from docx import Document
import openpyxl
import pandas as pd

from registrator_romania.new_request_registrator import prepare_users_data


class Transliterator:
    def __init__(self, language):
        if language == "tr":
            self.translit_dict = {
                "Ş": "S",
                "ş": "s",
                "İ": "I",
                "ı": "i",
                "Ğ": "G",
                "ğ": "g",
                "Ç": "C",
                "ç": "c",
                "Ö": "O",
                "ö": "o",
                "Ü": "U",
                "ü": "u",
            }
        else:
            self.translit_dict = {}

    def transliterate(self, text):
        return "".join(self.translit_dict.get(char, char) for char in text)


def get_users_data_from_docx():
    doc = Document("users.docx")
    user = False
    template = {
        "Nume Pasaport": "",
        "Prenume Pasaport": "",
        "Data nasterii": "",
        "Locul naşterii": "",
        "Prenume Mama": "",
        "Prenume Tata": "",
        "Adresa de email": "",
        "Serie și număr Pașaport": "",
    }
    users = []
    user = {}
    mapping = {
        "Prenume Pasaport": ["Prenume"],
        "Nume Pasaport": ["Nume"],
        "Data nasterii": ["Data naşterii"],
        "Locul naşterii": ["Locul naşterii"],
        "Prenume Mama": ["Prenumele mamei", "Numele mame", "Numele mamei"],
        "Prenume Tata": [
            "Prenumele tatalui",
            "Numele tatalui",
        ],
        "Adresa de email": ["Adresa de e-mail"],
        "Serie și număr Pașaport": ["Seria şi numar Paşaport"],
    }
    for paragraph in doc.paragraphs:
        text = paragraph.text.replace("\n", "").strip()
        if not text:
            continue

        record = re.findall(r"(^[\d\.]*)(.*)", text)[0][1]
        col, val = list(map(lambda v: v.strip(), record.split(":")))

        key = None
        for k, v in mapping.items():
            if col in v:
                key = k
                break

        assert key
        val = Transliterator("tr").transliterate(val)

        if key == "Data nasterii":
            try:
                dt = dateutil.parser.parse(val, dayfirst=True)
                # dt = datetime.strptime(val, "%Y-%m-%d")
            except Exception:
                dt = dateutil.parser.parse(val, dayfirst=False)
                # dt = datetime.strptime(val, "%d-%m-%Y")
            val = dt.strftime("%Y-%m-%d")
        elif key == "":
            val = val.lower()

        user[key] = val

        if key == "Serie și număr Pașaport":
            users.append(user.copy())
            user.clear()

    return prepare_users_data(users)


def get_users_data_from_csv():
    df = pd.read_csv("users.csv")
    users_data = df.to_dict("records")
    return prepare_users_data(users_data)


def get_users_data_from_txt():
    keys = [
        "Prenume Pasaport",
        "Nume Pasaport",
        "Data nasterii",
        "Locul naşterii",
        "Prenume Mama",
        "Prenume Tata",
        "Adresa de email",
        "Serie și număr Pașaport",
    ]
    with open("users.txt") as f:
        data = f.read()

    objs = []
    for values in data.split("\n\n"):
        obj = {}
        for k, v in zip(keys, values.split("\n")):
            v = v.split(":")[-1].strip()
            v = Transliterator("tr").transliterate(v)
            if k == "Data nasterii":
                try:
                    dt = datetime.strptime(v, "%Y-%m-%d")
                except Exception:
                    dt = datetime.strptime(v, "%d-%m-%Y")
                v = dt.strftime("%Y-%m-%d")

            obj[k] = v
        obj["Adresa de email"] = obj["Adresa de email"].lower()
        objs.append(obj)

    return prepare_users_data(objs)


def prepare_users_data(users_data: list[dict]):
    keys = [
        "Prenume Pasaport",
        "Nume Pasaport",
        "Data nasterii",
        "Locul naşterii",
        "Prenume Mama",
        "Prenume Tata",
        "Adresa de email",
        "Serie și număr Pașaport",
    ]
    objs = []
    for us_data in users_data:
        obj = {}
        for k, v in us_data.items():
            # Replace values like `Doğum tarihi:09.09.1976`
            v = v or ""
            v = v.split(":")[-1].strip()
            # Change turkey letters on english letters
            v = Transliterator("tr").transliterate(v)

            if k == "Data nasterii":
                try:
                    dt = dateutil.parser.parse(v, dayfirst=False)
                except dateutil.parser.ParserError:
                    dt = dateutil.parser.parse(v, dayfirst=True)

                # We need to format date like `1976-09-09`
                v = dt.strftime("%Y-%m-%d")
                assert datetime.strptime(v, "%Y-%m-%d")

            obj[k] = v

        # Tranform case
        obj["Nume Pasaport"] = obj["Nume Pasaport"].upper()
        obj["Prenume Pasaport"] = obj["Prenume Pasaport"].upper()
        obj["Adresa de email"] = obj["Adresa de email"].lower()
        objs.append(obj)

    assert all(k in obj for k in keys for obj in objs)
    return objs


def get_users_data_from_xslx():
    keys = [
        "Prenume Pasaport",
        "Nume Pasaport",
        "Data nasterii",
        "Locul naşterii",
        "Prenume Mama",
        "Prenume Tata",
        "Adresa de email",
        "Serie și număr Pașaport",
    ]

    w = openpyxl.load_workbook("users.xlsx")
    sheet = w.active
    data = []
    for row in sheet.iter_rows(min_row=0, max_row=None, values_only=True):
        if not data:
            assert row
            row = keys.copy()

        data.append(row)
    df = pd.DataFrame(data)
    df.columns = df.iloc[0]
    df.drop(df.index[0], inplace=True)
    objs_raw = df.to_dict("records")
    return prepare_users_data(objs_raw)
